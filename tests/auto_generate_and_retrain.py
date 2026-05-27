"""
Auto Generate Ratings -> Retrain -> Audit -> Excel Report
=========================================================
Script tu dong: Seed ratings cho user moi -> Retrain FunkSVD -> Kiem thu -> Xuat Excel
"""
import sys, os, time, json, random, subprocess
import numpy as np
import pandas as pd
from pathlib import Path
from datetime import datetime
from collections import Counter

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))
os.environ["PYTHONIOENCODING"] = "utf-8"

from src.database.db_config import DatabaseConnector
from config.settings import RATINGS_PATH, MOVIES_PATH, USERS_PATH, MODEL_OUTPUT_DIR

# ═══════════════════════════════════════════════════
# Cau hinh gu dien anh cho tung user
# ═══════════════════════════════════════════════════
USER_PROFILES = {
    7001: {"name": "Admin (Action/Sci-Fi Fan)", "genres": ["Action", "Sci-Fi", "Thriller"], "rating_range": (4.0, 5.0)},
    7011: {"name": "Admin2 (Drama/Romance Fan)", "genres": ["Drama", "Romance"], "rating_range": (4.5, 5.0)},
    7013: {"name": "User A (Animation/Fantasy Fan)", "genres": ["Animation", "Fantasy", "Adventure"], "rating_range": (4.0, 5.0)},
    7018: {"name": "User B (Horror/Thriller Fan)", "genres": ["Horror", "Thriller", "Mystery"], "rating_range": (4.0, 5.0)},
    7019: {"name": "RBAC Test (Comedy/Musical Fan)", "genres": ["Comedy", "Musical"], "rating_range": (3.5, 5.0)},
    7020: {"name": "Anime Fan (Animation/Action)", "genres": ["Animation", "Action", "Sci-Fi"], "rating_range": (4.5, 5.0)},
}

def section(t):
    print(f"\n{'='*65}\n  {t}\n{'='*65}")

# ═══════════════════════════════════════════════════
# GIAI DOAN 1: TU DONG SINH RATINGS
# ═══════════════════════════════════════════════════
def phase1_seed_ratings():
    section("GIAI DOAN 1: Tu dong sinh ratings cho User >= 7000")
    conn = DatabaseConnector.get_connection()
    cur = conn.cursor()

    # Lay danh sach phim moi (100 phim co movie_id lon nhat)
    cur.execute("""
        SELECT m.movie_id, m.genres_orig FROM movies m
        WHERE m.movie_id > 3883 ORDER BY m.movie_id DESC LIMIT 100
    """)
    new_movies = cur.fetchall()  # [(movie_id, genres_orig), ...]
    print(f"  Tim thay {len(new_movies)} phim moi (movie_id > 3883)")

    # Cung lay 200 phim cu de da dang hoa
    cur.execute("""
        SELECT m.movie_id, m.genres_orig FROM movies m
        WHERE m.movie_id <= 3883 ORDER BY RANDOM() LIMIT 200
    """)
    old_movies = cur.fetchall()

    all_candidate_movies = new_movies + old_movies
    print(f"  Tong phim ung vien: {len(all_candidate_movies)} (moi + cu)")

    # Lay user >= 7000
    cur.execute("SELECT user_id, email FROM users WHERE user_id >= 7000 ORDER BY user_id")
    users = cur.fetchall()
    print(f"  Tim thay {len(users)} user (>= 7000)")

    ratings_log = []  # Luu lai de xuat Excel
    total_inserted = 0

    for user_id, email in users:
        profile = USER_PROFILES.get(user_id)
        if not profile:
            profile = {"name": f"User {user_id}", "genres": ["Drama", "Action"], "rating_range": (3.0, 5.0)}

        preferred_genres = set(profile["genres"])
        lo, hi = profile["rating_range"]

        # Tim phim phu hop voi gu
        matched_movies = []
        other_movies = []
        for mid, genres_str in all_candidate_movies:
            if not genres_str:
                continue
            movie_genres = set(genres_str.split("|"))
            if movie_genres & preferred_genres:
                matched_movies.append((mid, genres_str))
            else:
                other_movies.append((mid, genres_str))

        # Chon 20-30 phim phu hop + 5-10 phim khac the loai
        n_matched = min(random.randint(20, 30), len(matched_movies))
        n_other = min(random.randint(5, 10), len(other_movies))

        selected = random.sample(matched_movies, n_matched) if n_matched > 0 else []
        selected += random.sample(other_movies, n_other) if n_other > 0 else []

        user_count = 0
        for mid, genres_str in selected:
            movie_genres = set(genres_str.split("|"))
            is_preferred = bool(movie_genres & preferred_genres)

            if is_preferred:
                rating = round(random.uniform(lo, hi) * 2) / 2  # Round to 0.5
            else:
                rating = round(random.uniform(1.0, 3.5) * 2) / 2

            rating = max(0.5, min(5.0, rating))
            ts = int(time.time()) - random.randint(0, 86400 * 30)

            try:
                cur.execute("""
                    INSERT INTO ratings (user_id, movie_id, rating, timestamp)
                    VALUES (%s, %s, %s, %s)
                    ON CONFLICT (user_id, movie_id) DO UPDATE
                    SET rating = EXCLUDED.rating, timestamp = EXCLUDED.timestamp
                """, (user_id, mid, rating, ts))
                user_count += 1
                total_inserted += 1

                ratings_log.append({
                    "user_id": user_id,
                    "user_name": profile["name"],
                    "movie_id": mid,
                    "genres": genres_str,
                    "rating": rating,
                    "is_preferred": is_preferred,
                    "is_new_movie": mid > 3883,
                })
            except Exception as e:
                pass

        print(f"    User {user_id} ({profile['name']}): {user_count} ratings "
              f"(gu: {', '.join(profile['genres'])})")

    conn.commit()
    cur.close(); conn.close()

    print(f"\n  [OK] Tong cong da chen: {total_inserted} ratings")
    print(f"  [OK] Cho {len(users)} users")
    return ratings_log

# ═══════════════════════════════════════════════════
# GIAI DOAN 2: RETRAIN MODEL
# ═══════════════════════════════════════════════════
def phase2_retrain():
    section("GIAI DOAN 2: Huan luyen lai mo hinh FunkSVD")

    from src.data.data_loader import DataLoader
    from src.data.preprocessor import DataPreprocessor
    from src.models.funk_svd import FunkSVDModel
    from src.evaluation.evaluator import Evaluator

    # 1. Load data tu file .dat (goc)
    loader = DataLoader(RATINGS_PATH, MOVIES_PATH, USERS_PATH)
    ratings_file = loader.load_ratings()
    print(f"  Ratings tu file .dat: {len(ratings_file)}")

    # 2. Load ratings tu DB (bao gom ratings moi)
    conn = DatabaseConnector.get_connection()
    cur = conn.cursor()
    cur.execute("SELECT user_id, movie_id, rating, timestamp FROM ratings")
    db_rows = cur.fetchall()
    cur.close(); conn.close()

    db_ratings = pd.DataFrame(db_rows, columns=["user_id", "movie_id", "rating", "timestamp"])
    print(f"  Ratings tu DB: {len(db_ratings)}")

    # 3. Merge: uu tien DB (co ratings moi), loai trung lap
    combined = pd.concat([ratings_file, db_ratings], ignore_index=True)
    combined = combined.drop_duplicates(subset=["user_id", "movie_id"], keep="last")
    combined = combined.sort_values("timestamp").reset_index(drop=True)
    print(f"  Ratings sau merge (loai trung): {len(combined)}")

    # 4. Map IDs TRUOC (de mapping khong bi anh huong boi oversampling)
    preprocessor = DataPreprocessor()
    ratings_mapped = preprocessor.map_ids(combined)

    # 5. Train/Test split
    train_df, test_df = preprocessor.time_based_split(ratings_mapped, test_ratio=0.2)

    # 5b. Dam bao ALL ratings cua user moi nam trong train (khong bi split vao test)
    #     De model hoc duoc latent vector cho ho -> su dung SVD thuan
    new_in_test = test_df[test_df["user_id"] >= 7000]
    if len(new_in_test) > 0:
        train_df = pd.concat([train_df, new_in_test], ignore_index=True)
        test_df = test_df[test_df["user_id"] < 7000].copy()
        print(f"  Chuyen {len(new_in_test)} ratings user moi tu Test -> Train")

    # 6. Oversampling: Nhan ban ratings cua user >= 7000 x3 (CHI trong train_df)
    new_user_train = train_df[train_df["user_id"] >= 7000].copy()
    n_new = len(new_user_train)
    if n_new > 0:
        oversampled = pd.concat([new_user_train] * 3, ignore_index=True)
        oversampled["rating"] = oversampled["rating"] + np.random.uniform(-0.05, 0.05, len(oversampled))
        oversampled["rating"] = oversampled["rating"].clip(0.5, 5.0)
        train_df = pd.concat([train_df, oversampled], ignore_index=True)
        print(f"  Oversampling: Nhan ban {n_new} x3 = +{len(oversampled)} ratings cho user moi")

    print(f"  Train: {len(train_df)} | Test: {len(test_df)}")

    # 7. Train FunkSVD voi epochs = 15
    model = FunkSVDModel(n_factors=50, n_epochs=15, lr=0.005, reg=0.02)

    epoch_log = []
    import io
    from contextlib import redirect_stdout

    # Capture epoch output
    old_stdout = sys.stdout
    captured = io.StringIO()

    start_time = time.time()
    # Temporarily redirect to capture epoch RMSE
    sys.stdout = captured
    model.fit(train_df)
    sys.stdout = old_stdout

    train_time = time.time() - start_time
    output = captured.getvalue()
    print(output)  # Print captured output

    # Parse epoch log
    for line in output.split("\n"):
        if "Epoch" in line and "RMSE" in line:
            parts = line.strip().split()
            try:
                epoch_num = int(parts[1].split("/")[0])
                rmse_val = float(parts[-1])
                epoch_log.append({"epoch": epoch_num, "train_rmse": rmse_val})
            except:
                pass

    print(f"\n  Thoi gian training: {train_time:.1f} giay")

    # 8. Evaluate — filter valid indices to avoid IndexError
    n_users_trained = model.user_factors.shape[0]
    n_items_trained = model.item_factors.shape[0]
    valid_test = test_df[
        (test_df["user_idx"] < n_users_trained) &
        (test_df["movie_idx"] < n_items_trained)
    ].copy()
    print(f"  Test samples (valid): {len(valid_test)}/{len(test_df)}")
    metrics = Evaluator.evaluate(model, valid_test, k=10)
    print(f"\n  === KET QUA DANH GIA ===")
    print(f"  RMSE:    {metrics['RMSE']}")
    print(f"  MAE:     {metrics['MAE']}")
    print(f"  NDCG@10: {metrics['NDCG@K']}")

    # 9. Save model
    model_path = MODEL_OUTPUT_DIR / "funksvd_model.pkl"
    model.save(model_path)
    print(f"  [OK] Model luu tai: {model_path}")

    # 10. Save mappings (de engine_wrapper dung dung index)
    import pickle
    mappings = {
        "user2idx": preprocessor.user2idx,
        "idx2user": preprocessor.idx2user,
        "movie2idx": preprocessor.movie2idx,
        "idx2movie": preprocessor.idx2movie,
    }
    mappings_path = MODEL_OUTPUT_DIR / "id_mappings.pkl"
    with open(mappings_path, "wb") as f:
        pickle.dump(mappings, f)
    print(f"  [OK] Mappings luu tai: {mappings_path}")
    print(f"       user2idx: {len(preprocessor.user2idx)} users")
    print(f"       movie2idx: {len(preprocessor.movie2idx)} movies")

    return epoch_log, metrics, train_time

# ═══════════════════════════════════════════════════
# GIAI DOAN 3: AUDIT GOI Y SAU RETRAIN
# ═══════════════════════════════════════════════════
def phase3_audit():
    section("GIAI DOAN 3: Kiem toan chat luong goi y (Post-Retrain)")

    import requests
    BASE = "http://127.0.0.1:8000"

    audit_results = []

    for user_id, profile in USER_PROFILES.items():
        preferred = set(profile["genres"])
        try:
            r = requests.get(f"{BASE}/recommend/{user_id}", timeout=10)
            if r.status_code != 200:
                print(f"  [FAIL] User {user_id}: HTTP {r.status_code}")
                audit_results.append({
                    "user_id": user_id, "user_name": profile["name"],
                    "strategy": "ERROR", "total_recs": 0,
                    "genre_match_pct": 0, "recommendations": [],
                })
                continue

            data = r.json()
            strategy = data.get("strategy", "unknown")
            recs = data.get("recommendations", [])

            # Tinh ty le khop the loai
            match_count = 0
            rec_details = []
            for i, rec in enumerate(recs[:10]):
                genres_str = rec.get("genres_orig", "")
                movie_genres = set(genres_str.split("|")) if genres_str else set()
                is_match = bool(movie_genres & preferred)
                if is_match:
                    match_count += 1
                rec_details.append({
                    "rank": i + 1,
                    "movie_id": rec.get("movie_id"),
                    "title": rec.get("title", "N/A"),
                    "genres": genres_str,
                    "predicted_rating": rec.get("predicted_rating"),
                    "genre_match": "CO" if is_match else "KHONG",
                })

            match_pct = (match_count / min(len(recs), 10)) * 100 if recs else 0
            is_svd = strategy in ("hybrid", "hybrid_foldin")

            print(f"  User {user_id} ({profile['name']}):")
            print(f"    Strategy: {strategy} | Khop gu: {match_count}/10 ({match_pct:.0f}%) "
                  f"| SVD: {'CO' if is_svd else 'KHONG'}")

            audit_results.append({
                "user_id": user_id,
                "user_name": profile["name"],
                "preferred_genres": ", ".join(profile["genres"]),
                "strategy": strategy,
                "is_svd": is_svd,
                "total_recs": len(recs),
                "genre_match_count": match_count,
                "genre_match_pct": match_pct,
                "recommendations": rec_details,
            })
        except Exception as e:
            print(f"  [ERROR] User {user_id}: {e}")
            audit_results.append({
                "user_id": user_id, "user_name": profile["name"],
                "strategy": "ERROR", "total_recs": 0,
                "genre_match_pct": 0, "recommendations": [],
            })

    return audit_results

# ═══════════════════════════════════════════════════
# GIAI DOAN 4: XUAT EXCEL
# ═══════════════════════════════════════════════════
def phase4_export_excel(ratings_log, epoch_log, metrics, train_time, audit_results):
    section("GIAI DOAN 4: Xuat bao cao Excel")

    try:
        import openpyxl
    except ImportError:
        subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl"], capture_output=True)
        import openpyxl

    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    hdr_font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    hdr_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    title_font = Font(name="Arial", bold=True, size=14, color="1F4E79")
    bold_font = Font(name="Arial", bold=True, size=10)
    norm_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center")
    border = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))

    def style_header(ws, row, cols):
        for c in range(1, cols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = hdr_font; cell.fill = hdr_fill
            cell.alignment = center; cell.border = border

    # ── Sheet 1: Nhat Ky Tao Ratings ──
    ws1 = wb.active
    ws1.title = "Nhat Ky Tao Ratings"
    ws1.merge_cells("A1:G1")
    ws1["A1"] = "NHAT KY TAO RATINGS GIA LAP CHO USER MOI (>= 7000)"
    ws1["A1"].font = title_font
    ws1["A2"] = f"Ngay: {datetime.now().strftime('%Y-%m-%d %H:%M')}"

    headers1 = ["STT", "User ID", "Ten User", "Movie ID", "The Loai", "Diem", "Phu Hop Gu?"]
    for c, h in enumerate(headers1, 1):
        ws1.cell(row=4, column=c, value=h)
    style_header(ws1, 4, 7)

    for i, r in enumerate(ratings_log):
        row = 5 + i
        ws1.cell(row=row, column=1, value=i+1).alignment = center
        ws1.cell(row=row, column=2, value=r["user_id"]).alignment = center
        ws1.cell(row=row, column=3, value=r["user_name"])
        ws1.cell(row=row, column=4, value=r["movie_id"]).alignment = center
        ws1.cell(row=row, column=5, value=r["genres"])
        ws1.cell(row=row, column=6, value=r["rating"]).alignment = center
        cell = ws1.cell(row=row, column=7, value="CO" if r["is_preferred"] else "KHONG")
        cell.alignment = center
        cell.fill = pass_fill if r["is_preferred"] else fail_fill
        for c in range(1, 8):
            ws1.cell(row=row, column=c).border = border

    # Thong ke cuoi sheet
    summary_row = 5 + len(ratings_log) + 2
    ws1.cell(row=summary_row, column=1, value="TONG KET:").font = bold_font
    ws1.cell(row=summary_row+1, column=1, value=f"Tong ratings da tao: {len(ratings_log)}")
    n_preferred = sum(1 for r in ratings_log if r["is_preferred"])
    ws1.cell(row=summary_row+2, column=1, value=f"Ratings phu hop gu: {n_preferred}/{len(ratings_log)}")
    n_new = sum(1 for r in ratings_log if r["is_new_movie"])
    ws1.cell(row=summary_row+3, column=1, value=f"Ratings tren phim moi: {n_new}")

    for w, width in [(1,6),(2,10),(3,30),(4,10),(5,40),(6,8),(7,14)]:
        ws1.column_dimensions[chr(64+w)].width = width

    # ── Sheet 2: Tien Trinh Huan Luyen ──
    ws2 = wb.create_sheet("Tien Trinh Huan Luyen")
    ws2.merge_cells("A1:D1")
    ws2["A1"] = "TIEN TRINH HUAN LUYEN LAI MO HINH FUNKSVD"
    ws2["A1"].font = title_font
    ws2["A2"] = f"Cau hinh: 50 factors, 15 epochs, lr=0.005, reg=0.02, oversampling x3"

    headers2 = ["Epoch", "Train RMSE", "Bien Dong", "Ghi Chu"]
    for c, h in enumerate(headers2, 1):
        ws2.cell(row=4, column=c, value=h)
    style_header(ws2, 4, 4)

    for i, ep in enumerate(epoch_log):
        row = 5 + i
        ws2.cell(row=row, column=1, value=ep["epoch"]).alignment = center
        ws2.cell(row=row, column=2, value=round(ep["train_rmse"], 4)).alignment = center
        if i > 0:
            delta = ep["train_rmse"] - epoch_log[i-1]["train_rmse"]
            ws2.cell(row=row, column=3, value=round(delta, 4)).alignment = center
        else:
            ws2.cell(row=row, column=3, value="-").alignment = center
        note = ""
        if i == 0: note = "Khoi tao"
        elif i == len(epoch_log)-1: note = "Hoi tu"
        elif ep["train_rmse"] < 1.0: note = "RMSE < 1.0"
        ws2.cell(row=row, column=4, value=note)
        for c in range(1, 5):
            ws2.cell(row=row, column=c).border = border

    # Metrics summary
    mr = 5 + len(epoch_log) + 2
    ws2.cell(row=mr, column=1, value="CHI SO DANH GIA CUOI CUNG:").font = bold_font
    for i, (k, v) in enumerate([("Test RMSE", metrics["RMSE"]), ("Test MAE", metrics["MAE"]),
                                  ("NDCG@10", metrics["NDCG@K"]), ("Thoi gian train", f"{train_time:.1f}s")]):
        ws2.cell(row=mr+1+i, column=1, value=k).font = bold_font
        ws2.cell(row=mr+1+i, column=2, value=v)

    for w, width in [(1,10),(2,14),(3,12),(4,20)]:
        ws2.column_dimensions[chr(64+w)].width = width

    # ── Sheet 3: Doi Chieu Goi Y ──
    ws3 = wb.create_sheet("Doi Chieu Goi Y")
    ws3.merge_cells("A1:H1")
    ws3["A1"] = "DOI CHIEU GOI Y HAU HUAN LUYEN (POST-RETRAIN AUDIT)"
    ws3["A1"].font = title_font

    current_row = 3
    for audit in audit_results:
        uid = audit["user_id"]
        # User header
        ws3.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=8)
        cell = ws3.cell(row=current_row, column=1,
                        value=f"User {uid} - {audit['user_name']} | "
                              f"Strategy: {audit['strategy']} | "
                              f"Gu: {audit.get('preferred_genres','')} | "
                              f"Khop: {audit.get('genre_match_count',0)}/10 ({audit.get('genre_match_pct',0):.0f}%)")
        cell.font = Font(name="Arial", bold=True, size=11, color="2E75B6")
        current_row += 1

        # Table header
        headers3 = ["Hang", "Movie ID", "Ten Phim", "The Loai", "Diem Du Doan", "Khop Gu?", "Phim Moi?", "SVD?"]
        for c, h in enumerate(headers3, 1):
            ws3.cell(row=current_row, column=c, value=h)
        style_header(ws3, current_row, 8)
        current_row += 1

        for rec in audit.get("recommendations", []):
            ws3.cell(row=current_row, column=1, value=rec["rank"]).alignment = center
            ws3.cell(row=current_row, column=2, value=rec["movie_id"]).alignment = center
            ws3.cell(row=current_row, column=3, value=rec["title"])
            ws3.cell(row=current_row, column=4, value=rec["genres"])
            ws3.cell(row=current_row, column=5, value=rec.get("predicted_rating")).alignment = center
            match_cell = ws3.cell(row=current_row, column=6, value=rec["genre_match"])
            match_cell.alignment = center
            match_cell.fill = pass_fill if rec["genre_match"] == "CO" else fail_fill
            is_new = rec["movie_id"] > 3883 if rec["movie_id"] else False
            ws3.cell(row=current_row, column=7, value="CO" if is_new else "KHONG").alignment = center
            ws3.cell(row=current_row, column=8, value="CO" if audit.get("is_svd") else "KHONG").alignment = center
            for c in range(1, 9):
                ws3.cell(row=current_row, column=c).border = border
            current_row += 1

        current_row += 1  # Empty row between users

    for col, w in [("A",6),("B",10),("C",45),("D",35),("E",12),("F",10),("G",10),("H",8)]:
        ws3.column_dimensions[col].width = w

    # Save
    out_path = BASE_DIR / "output" / "BaoCao_KiemThu_Retrain_UserMoi.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    print(f"\n  [OK] Bao cao Excel: {out_path}")
    return out_path


# ═══════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════
if __name__ == "__main__":
    print("=" * 65)
    print("  AUTO GENERATE RATINGS -> RETRAIN -> AUDIT -> EXCEL")
    print(f"  {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 65)

    start = time.time()

    # Phase 1
    ratings_log = phase1_seed_ratings()

    # Phase 2
    epoch_log, metrics, train_time = phase2_retrain()

    # Phase 3 (can backend dang chay)
    print("\n  [INFO] Kiem tra backend tai http://127.0.0.1:8000 ...")
    try:
        import requests
        r = requests.get("http://127.0.0.1:8000/status", timeout=3)
        backend_ok = True
    except:
        backend_ok = False

    if backend_ok:
        audit_results = phase3_audit()
    else:
        print("  [WARN] Backend chua chay! Bo qua Phase 3 audit API.")
        print("  [WARN] De chay Phase 3, hay bat backend truoc.")
        audit_results = []
        for uid, prof in USER_PROFILES.items():
            audit_results.append({
                "user_id": uid, "user_name": prof["name"],
                "preferred_genres": ", ".join(prof["genres"]),
                "strategy": "CHUA KIEM TRA (backend off)",
                "is_svd": False, "total_recs": 0,
                "genre_match_count": 0, "genre_match_pct": 0,
                "recommendations": [],
            })

    # Phase 4
    report_path = phase4_export_excel(ratings_log, epoch_log, metrics, train_time, audit_results)

    elapsed = time.time() - start
    section(f"HOAN TAT trong {elapsed:.1f}s")
    print(f"  Excel: {report_path}")
    print(f"  Ratings da tao: {len(ratings_log)}")
    print(f"  RMSE cuoi: {metrics['RMSE']} | NDCG@10: {metrics['NDCG@K']}")
